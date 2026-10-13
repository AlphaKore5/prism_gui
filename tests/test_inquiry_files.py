"""
[3] Proper formatting and saving of files.

  a. the quotation a customer receives is a PDF on the company's name, with
     ₹ in lakh grouping and an aligned table — not a bare CSV and a
     fixed-width text table that falls apart in Gmail;
  b. the register can be saved as an .xlsx that reads the same on every
     locale (real dates, phone numbers keep their leading 0);
  c. a saved attachment's name starts with its inquiry number, keeps the
     sender's own name and extension, and never overwrites another;
  d. the mail itself is kept in the inquiry's folder as an .eml.
"""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PySide6.QtWidgets import QApplication  # noqa: E402

import core_bridge as CB  # noqa: E402
from core import inbox, mailflow, quoting, register, triage  # noqa: E402
from addons.inquiry import autoquote, documents  # noqa: E402

_app = QApplication.instance() or QApplication([])


def a_quote(**terms):
    return quoting.Quotation(
        number="QTN/26-27/0001", date=date(2026, 10, 13),
        customer="KJ <Pharmatech> & Co", contact="Mr Shah", email="buyer@kj.in",
        inquiry_no="INQ/26-27/0008",
        lines=[quoting.QuoteLine("SS Locker 8 Comp", Decimal("10"), "nos",
                                 Decimal("78000"), "9403")],
        terms=quoting.Terms(gst_percent=Decimal(18), validity_days=15,
                            payment="50% advance", delivery="3 weeks", **terms))


def mail(subject="Enquiry for lockers", body="Please quote 10 nos SS Locker 8 Comp.",
         message_id="<f1@x>", attachments=()):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = "Buyer <buyer@kj.in>"
    msg["To"] = "sales@acme.co.in"
    msg["Date"] = "Mon, 12 Oct 2026 09:14:00 +0530"
    msg["Message-ID"] = message_id
    msg.set_content(body)
    for name, data, mime in attachments:
        maintype, _, subtype = mime.partition("/")
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    return inbox.parse_message(msg.as_bytes(), uid=1)


class TheQuotationPDF(unittest.TestCase):

    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.settings = {"company": "Acme Springs",
                         "signature": "Ravi Patel\n+91 98765 43210"}

    def _text(self, path):
        from pypdf import PdfReader
        return " ".join(" ".join(page.extract_text().split())
                        for page in PdfReader(path).pages)

    def test_both_files_are_written_beside_each_other(self):
        csv_path, pdf_path = documents.quotation_files(a_quote(), self.settings,
                                                       self.folder)
        self.assertEqual(Path(csv_path).name, "QTN-26-27-0001.csv")
        self.assertEqual(Path(pdf_path).name, "QTN-26-27-0001.pdf")
        with open(pdf_path, "rb") as f:
            self.assertEqual(f.read(5), b"%PDF-")

    def test_it_reads_as_a_quotation_on_the_companys_name(self):
        _, pdf_path = documents.quotation_files(a_quote(), self.settings, self.folder)
        text = self._text(pdf_path)
        for expected in ("Acme Springs", "QTN/26-27/0001", "INQ/26-27/0008",
                         "SS Locker 8 Comp", "7,80,000.00", "9,20,400.00",
                         "3 weeks", "50% advance", "Ravi Patel", "13-10-2026"):
            self.assertIn(expected, text, expected)

    def test_customer_text_is_escaped_not_rendered(self):
        page = documents.quotation_html(a_quote(), self.settings)
        self.assertIn("KJ &lt;Pharmatech&gt; &amp; Co", page)
        self.assertNotIn("<Pharmatech>", page)

    def test_discount_and_freight_appear_when_set(self):
        page = documents.quotation_html(
            a_quote(discount_percent=Decimal(5), freight=Decimal(1500)),
            self.settings)
        self.assertIn("Discount 5%", page)
        self.assertIn("Freight", page)

    def test_the_customer_is_sent_the_pdf(self):
        csv_path, pdf_path = documents.quotation_files(a_quote(), self.settings,
                                                       self.folder)
        sent = documents.attachment_for(csv_path, pdf_path)
        self.assertEqual(sent["mime"], "application/pdf")
        self.assertEqual(sent["name"], "QTN-26-27-0001.pdf")

    def test_the_csv_goes_if_the_pdf_could_not_be_drawn(self):
        sent = documents.attachment_for("/x/QTN-26-27-0001.csv", "")
        self.assertEqual(sent["mime"], "text/csv")

    def test_the_mail_body_is_not_a_fixed_width_table(self):
        subject, body = autoquote.covering_mail(a_quote(), self.settings)
        self.assertIn("1. SS Locker 8 Comp — 10 nos × ₹ 78,000.00 = ₹ 7,80,000.00",
                      body)
        self.assertIn("Total: ₹ 9,20,400.00", body)
        self.assertNotIn("  " * 6, body)

    def test_our_own_pdf_is_never_taken_for_a_drawing(self):
        self.assertTrue(documents.is_ours("QTN-26-27-0001.pdf"))
        self.assertFalse(documents.is_ours("INQ-26-27-0008_drawing.pdf"))
        self.assertFalse(documents.is_ours("drawing.pdf"))

    def test_a_customers_own_po_is_not_taken_for_ours(self):
        """Review F11: a PO saved before attachments carried the INQ- prefix."""
        for name in ("PO-26-27-0012.pdf", "SO-25-26-7.xlsx", "INV-26-27-0001.csv"):
            self.assertFalse(documents.is_ours(name), name)
        for name in ("QTN-26-27-0001.csv", "INQ-26-27-0005_mail_2026-10-12_0914.eml",
                     "INQ-26-27-0005_mail_2026-10-12_0914-2.eml",
                     "mail_mail_undated.eml"):
            self.assertTrue(documents.is_ours(name), name)

    def test_a_pdf_that_cannot_be_written_is_never_reported_with_an_old_one(self):
        """Review F6: on Windows the PDF open in Acrobat cannot be replaced.
        QPdfWriter only logs that, and the old check passed on the stale
        file of the same name, which then went to the customer."""
        target = Path(self.folder) / "QTN-26-27-0001.pdf"
        target.write_bytes(b"%PDF-OLD stale quotation")

        def locked(*_a, **_k):
            raise PermissionError(13, "The process cannot access the file")
        real = os.replace
        os.replace = locked
        try:
            written = documents.quotation_pdf(a_quote(), self.settings, target)
        finally:
            os.replace = real
        self.assertEqual(written, "")
        self.assertEqual(target.read_bytes(), b"%PDF-OLD stale quotation")
        self.assertEqual([p.name for p in Path(self.folder).iterdir()],
                         ["QTN-26-27-0001.pdf"])          # no .part left behind
        self.assertIn("CSV", documents.pdf_fallback_note())

    def test_a_folder_that_cannot_be_written_returns_nothing(self):
        if os.name == "nt" or (hasattr(os, "geteuid") and os.geteuid() == 0):
            self.skipTest("needs POSIX permissions and a non-root user")
        target = Path(self.folder) / "QTN-26-27-0001.pdf"
        target.write_bytes(b"%PDF-OLD stale quotation")
        os.chmod(self.folder, 0o555)
        try:
            written = documents.quotation_pdf(a_quote(), self.settings, target)
        finally:
            os.chmod(self.folder, 0o755)
        self.assertEqual(written, "")
        self.assertEqual(target.read_bytes(), b"%PDF-OLD stale quotation")

    def test_a_new_pdf_replaces_an_old_one(self):
        target = Path(self.folder) / "QTN-26-27-0001.pdf"
        target.write_bytes(b"%PDF-OLD stale quotation")
        written = documents.quotation_pdf(a_quote(), self.settings, target)
        self.assertEqual(written, str(target))
        self.assertIn("Acme Springs", self._text(written))


class TheRegisterForExcel(unittest.TestCase):

    def test_dates_phones_and_money_keep_their_meaning(self):
        from openpyxl import load_workbook
        row = register.blank_row()
        row.update({"Inquiry no": "INQ/26-27/0001", "Date received": "02-08-2026",
                    "Phone": "09876543210", "Status": register.QUOTED,
                    "Quotation value": "142500.00", "Thread": "<a@x>",
                    "Salesman": "Ravi"})
        path = documents.register_xlsx([row], Path(tempfile.mkdtemp()) / "reg.xlsx")
        sheet = load_workbook(path).active
        header = [c.value for c in sheet[1]]
        values = dict(zip(header, [c for c in sheet[2]]))
        self.assertNotIn("Thread", header)
        self.assertIn("Salesman", header)
        received = values["Date received"]
        self.assertEqual(received.value.date(), date(2026, 8, 2))
        self.assertEqual(received.number_format, "DD-MM-YYYY")
        self.assertEqual(values["Phone"].value, "09876543210")
        self.assertEqual(values["Quotation value"].value, 142500.0)
        self.assertEqual(sheet.freeze_panes, "B2")
        self.assertTrue(values["Status"].fill.fgColor.rgb.startswith("FF"))

    def test_an_unreadable_date_is_kept_as_typed(self):
        from openpyxl import load_workbook
        row = {"Inquiry no": "X", "Date received": "early Aug"}
        path = documents.register_xlsx([row], Path(tempfile.mkdtemp()) / "r.xlsx")
        sheet = load_workbook(path).active
        header = [c.value for c in sheet[1]]
        self.assertEqual(sheet.cell(2, header.index("Date received") + 1).value,
                         "early Aug")


class CustomerTextInExcelIsOnlyText(unittest.TestCase):
    """Review F4 and F5: the register's text comes from other people's mail."""

    def _sheet(self, **fields):
        from openpyxl import load_workbook
        row = dict(register.blank_row(), **{"Inquiry no": "INQ/26-27/0001"}, **fields)
        path = documents.register_xlsx([row], Path(tempfile.mkdtemp()) / "r.xlsx")
        sheet = load_workbook(path).active
        header = [c.value for c in sheet[1]]
        return {h: c for h, c in zip(header, sheet[2])}

    def test_a_formula_from_a_customer_is_stored_as_text(self):
        cells = self._sheet(**{"Customer": '=HYPERLINK("http://evil.example","Click")',
                               "Product asked": "=1+1", "Notes": "@SUM(A1)",
                               "Contact person": "+91 98765", "Quantity": "-5 nos"})
        for column in ("Customer", "Product asked", "Notes", "Contact person",
                       "Quantity"):
            self.assertEqual(cells[column].data_type, "s", column)
        self.assertEqual(cells["Product asked"].value, "=1+1")

    def test_control_characters_do_not_stop_the_export(self):
        cells = self._sheet(Notes="OCR\x01text\x07 read", Customer="A\x00B")
        self.assertEqual(cells["Notes"].value, "OCRtext read")
        self.assertEqual(cells["Customer"].value, "AB")


class NamesThatWorkOnWindowsAndMac(unittest.TestCase):

    def test_reserved_characters_and_names(self):
        self.assertEqual(documents.safe_name('a/b\\c:d*e?f"g<h>i|j'),
                         "a-b-c-d-e-f-g-h-i-j")
        self.assertEqual(documents.safe_name("CON"), "_CON")
        self.assertEqual(documents.safe_name("report. "), "report")
        self.assertLessEqual(len(documents.safe_name("x" * 500)), 80)


class _Check(unittest.TestCase):

    def setUp(self):
        self.root = tempfile.mkdtemp()
        self.paths = mailflow.Paths(self.root)

    def check(self, messages):
        real = inbox.fetch_new
        inbox.fetch_new = lambda cfg, state, **_kw: (messages, inbox.State(1, 99), "")
        try:
            return mailflow.check({"api_key": "", "inbox": {}}, self.paths,
                                  knowledge=triage.Knowledge(customers={"kj.in"}),
                                  local_only=True)
        finally:
            inbox.fetch_new = real


class AttachmentsCarryTheirInquiryNumber(_Check):

    def test_the_name_starts_with_the_number_and_keeps_the_rest(self):
        result = self.check([mail(attachments=[
            ("3F2A9C.jpeg", b"\xff\xd8\xff", "image/jpeg"),
            ("drawing.pdf", b"%PDF-1.4", "application/pdf")])])
        names = sorted(Path(f).name for f in result.new_inquiries[0].files)
        self.assertEqual(names, ["INQ-26-27-0001_3F2A9C.jpeg",
                                 "INQ-26-27-0001_drawing.pdf"])
        row = register.load(self.paths.register_csv)[0]
        self.assertIn("INQ-26-27-0001_drawing.pdf", row["Drawing"])

    def test_two_files_with_one_name_both_survive(self):
        self.check([mail(attachments=[("drawing.pdf", b"%PDF-1", "application/pdf")])])
        result = self.check([mail("Re: Enquiry for lockers", "revised",
                                  message_id="<f2@x>",
                                  attachments=[("drawing.pdf", b"%PDF-2",
                                                "application/pdf")])])
        names = sorted(p.name for p in Path(result.replies[0].folder).iterdir()
                       if p.suffix == ".pdf")
        self.assertEqual(names, ["INQ-26-27-0001_drawing-2.pdf",
                                 "INQ-26-27-0001_drawing.pdf"])

    def test_a_hostile_name_still_cannot_leave_the_folder(self):
        result = self.check([mail(attachments=[
            ("../../evil.pdf", b"%PDF", "application/pdf")])])
        path = Path(result.new_inquiries[0].files[0])
        self.assertEqual(path.parent, Path(result.new_inquiries[0].folder))

    def test_without_a_prefix_nothing_changes(self):
        folder = tempfile.mkdtemp()
        written = inbox.save_attachments(
            mail(attachments=[("a.pdf", b"x", "application/pdf")]), folder)
        self.assertEqual(Path(written[0]).name, "a.pdf")


class TheMailItselfIsKept(_Check):

    def test_an_enquiry_and_its_reply_are_both_saved_as_eml(self):
        self.check([mail()])
        result = self.check([mail("Re: Enquiry for lockers", "10 more please",
                                  message_id="<f3@x>")])
        folder = Path(result.replies[0].folder)
        saved = sorted(p.name for p in folder.glob("*.eml"))
        self.assertEqual(len(saved), 2)
        self.assertTrue(all(n.startswith("INQ-26-27-0001_mail_") for n in saved))
        subjects = sorted(inbox.parse_message((folder / n).read_bytes()).subject
                          for n in saved)
        self.assertEqual(subjects, ["Enquiry for lockers", "Re: Enquiry for lockers"])

    def test_the_eml_is_not_a_drawing(self):
        self.check([mail()])
        folder = Path(register.load(self.paths.register_csv)[0]["Folder"])
        self.assertFalse(any(p.suffix.lower() in (".pdf", ".dwg", ".dxf")
                             for p in folder.iterdir()))


if __name__ == "__main__":
    unittest.main()
