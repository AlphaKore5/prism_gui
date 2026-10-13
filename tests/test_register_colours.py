"""The inquiry register's colours ([12][15][17], the owner's decision of
13 Oct 2026).

One colour per state, the same on every register tab and in the Excel copy:
yet to quote blue, quoted grey, no reply amber, negotiating orange, won
green, lost red. Before this, New, Accepted and Converted were one green and
Quoted and Following up two greys nobody could tell apart, and a quoted row
vanished from "To quote" the moment it was quoted.

What these defend:
  · each status has its owner-given colour family, in light and in dark;
  · every two states are clearly apart, and the text on each is readable;
  · a picked row still reads as picked on every tint;
  · a quoted row stays on To quote, grey, under the ones still to quote,
    and is never counted as "to quote" anywhere;
  · the .xlsx carries the same colours row by row, with a colour key.
"""
from __future__ import annotations

import colorsys
import itertools
import math
import os
import sys
import tempfile
import unittest
from email.message import EmailMessage
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PySide6.QtWidgets import QApplication  # noqa: E402

import core_bridge as CB  # noqa: E402
import dashboard_data as DATA  # noqa: E402
import theme  # noqa: E402
from addons.inquiry import dialog as UI  # noqa: E402
from addons.inquiry import documents  # noqa: E402
from core import inbox, mailflow, register  # noqa: E402

_app = QApplication.instance() or QApplication([])

_REAL_SAVE = CB.config.save
_REAL_CONFIG_PATH = CB.config.CONFIG_PATH
_SCRATCH = tempfile.mkdtemp(prefix="prism-test-config-")


def _refuse(cfg):
    raise AssertionError("config.save() called from a colour test")


def setUpModule():
    CB.config.save = _refuse
    CB.config.CONFIG_PATH = os.path.join(_SCRATCH, "config.json")


def tearDownModule():
    CB.config.save = _REAL_SAVE
    CB.config.CONFIG_PATH = _REAL_CONFIG_PATH


# ── colour arithmetic (WCAG 2.x contrast, CIEDE2000) ───────────────────────
def _rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))


def _linear(channel: float) -> float:
    channel /= 255
    return channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4


def _luminance(value) -> float:
    r, g, b = (_linear(c) for c in (_rgb(value) if isinstance(value, str) else value))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b) -> float:
    hi, lo = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def _lab(value):
    r, g, b = (_linear(c) for c in (_rgb(value) if isinstance(value, str) else value))
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = 0.2126 * r + 0.7152 * g + 0.0722 * b
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883
    f = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116  # noqa: E731
    return 116 * f(y) - 16, 500 * (f(x) - f(y)), 200 * (f(y) - f(z))


def delta_e(a, b) -> float:
    """CIEDE2000. About 2 is the smallest step a person notices side by
    side; 10 and up reads as a different colour."""
    L1, a1, b1 = _lab(a)
    L2, a2, b2 = _lab(b)
    cbar = (math.hypot(a1, b1) + math.hypot(a2, b2)) / 2
    g = 0.5 * (1 - math.sqrt(cbar ** 7 / (cbar ** 7 + 25 ** 7)))
    a1p, a2p = (1 + g) * a1, (1 + g) * a2
    c1p, c2p = math.hypot(a1p, b1), math.hypot(a2p, b2)
    h1p = math.degrees(math.atan2(b1, a1p)) % 360
    h2p = math.degrees(math.atan2(b2, a2p)) % 360
    dlp, dcp = L2 - L1, c2p - c1p
    dh = 0.0 if c1p * c2p == 0 else h2p - h1p
    if dh > 180:
        dh -= 360
    elif dh < -180:
        dh += 360
    dhp = 2 * math.sqrt(c1p * c2p) * math.sin(math.radians(dh / 2))
    lbar, cbarp = (L1 + L2) / 2, (c1p + c2p) / 2
    if c1p * c2p == 0:
        hbar = h1p + h2p
    elif abs(h1p - h2p) <= 180:
        hbar = (h1p + h2p) / 2
    else:
        hbar = (h1p + h2p + (360 if h1p + h2p < 360 else -360)) / 2
    t = (1 - 0.17 * math.cos(math.radians(hbar - 30))
         + 0.24 * math.cos(math.radians(2 * hbar))
         + 0.32 * math.cos(math.radians(3 * hbar + 6))
         - 0.20 * math.cos(math.radians(4 * hbar - 63)))
    dtheta = 30 * math.exp(-((hbar - 275) / 25) ** 2)
    rc = 2 * math.sqrt(cbarp ** 7 / (cbarp ** 7 + 25 ** 7))
    sl = 1 + 0.015 * (lbar - 50) ** 2 / math.sqrt(20 + (lbar - 50) ** 2)
    sc, sh = 1 + 0.045 * cbarp, 1 + 0.015 * cbarp * t
    rt = -math.sin(math.radians(2 * dtheta)) * rc
    return math.sqrt((dlp / sl) ** 2 + (dcp / sc) ** 2 + (dhp / sh) ** 2
                     + rt * (dcp / sc) * (dhp / sh))


def _hue_and_saturation(value: str) -> tuple[float, float]:
    h, _l, s = colorsys.rgb_to_hls(*(c / 255 for c in _rgb(value)))
    return h * 360, s


# The owner's words, as hue ranges (degrees, HSL). "quoted" is the grey.
FAMILIES = {
    "new": (200, 235),          # blue
    "no_reply": (38, 56),       # amber
    "negotiating": (15, 37),    # orange
    "won": (115, 165),          # green
    "lost": (-12, 8),           # red (wraps through 0)
}
MODES = ("light", "dark")
CANVAS = {"light": "#ffffff", "dark": "#18181b"}


class EachStateHasItsOwnColour(unittest.TestCase):

    def test_the_owners_mapping(self):
        expected = {register.NEW: "new", register.QUOTED: "quoted",
                    register.FOLLOWING_UP: "no_reply",
                    register.NEGOTIATING: "negotiating",
                    register.ACCEPTED: "won", register.CONVERTED: "won",
                    register.NOT_CONVERTED: "lost"}
        self.assertEqual(UI.STATUS_STATES, expected)
        for status in register.STATUSES:
            self.assertIn(status, UI.STATUS_COLOURS, status)

    def test_every_state_has_a_light_and_a_dark_pair(self):
        for mode in MODES:
            self.assertEqual(set(theme.REGISTER_TONES[mode]),
                             set(theme.REGISTER_STATES), mode)

    def test_each_state_is_in_its_colour_family_in_both_themes(self):
        for mode in MODES:
            for state, (bg, _ink) in theme.REGISTER_TONES[mode].items():
                hue, saturation = _hue_and_saturation(bg)
                if state == "quoted":
                    self.assertLess(saturation, 0.12, f"{mode} {state} {bg}")
                    continue
                low, high = FAMILIES[state]
                hue = hue - 360 if hue > 180 and low < 0 else hue
                self.assertTrue(low <= hue <= high,
                                f"{mode} {state} {bg}: hue {hue:.0f}")
                self.assertGreater(saturation, 0.4, f"{mode} {state} {bg}")

    def test_every_two_states_are_clearly_apart_in_both_themes(self):
        """10 in light, where the tints are soft washes (the look the owner
        is used to) and each also has its own coloured ink; 15 in dark."""
        floor = {"light": 10, "dark": 15}
        for mode in MODES:
            tones = theme.REGISTER_TONES[mode]
            for a, b in itertools.combinations(tones, 2):
                d = delta_e(tones[a][0], tones[b][0])
                self.assertGreaterEqual(d, floor[mode],
                                        f"{mode}: {a} vs {b} ΔE00 {d:.1f}")

    def test_the_tints_are_soft_washes_in_light(self):
        """Round 2's strong tints (#a9cdfc blue, #fdb071 orange) read as
        selected rows. Every light tint is a pale wash: lighter than the
        old blue, the way the baseline's greens were."""
        for state, (bg, _ink) in theme.REGISTER_TONES["light"].items():
            # relative luminance; the old blue was 0.59, white is 1
            self.assertGreaterEqual(_luminance(bg), 0.7, f"{state} {bg}")

    def test_the_text_on_each_is_readable(self):
        """7:1 (WCAG AAA) on the light set, which is what Prism paints;
        4.5:1 (AA) on the dark set, where white on a true orange cannot
        reach 7:1 without turning it brown."""
        floor = {"light": 7.0, "dark": 4.5}
        for mode in MODES:
            for state, (bg, ink) in theme.REGISTER_TONES[mode].items():
                self.assertGreaterEqual(contrast(bg, ink), floor[mode],
                                        f"{mode} {state} {bg}/{ink}")

    def test_each_tint_shows_against_its_page(self):
        for mode in MODES:
            for state, (bg, _ink) in theme.REGISTER_TONES[mode].items():
                self.assertGreaterEqual(delta_e(bg, CANVAS[mode]), 5,
                                        f"{mode} {state} {bg}")

    def test_no_role_accent_rotation_can_move_them(self):
        for mode in MODES:
            for bg, ink in theme.REGISTER_TONES[mode].values():
                self.assertNotIn(bg.lower(), theme._ACCENT_HEXES)
                self.assertNotIn(ink.lower(), theme._ACCENT_HEXES)

    def test_status_colours_follows_the_palette(self):
        dark = UI.status_colours(dark=True)
        self.assertEqual(dark[register.NEW], theme.REGISTER_TONES["dark"]["new"])
        self.assertEqual(UI.status_colours(dark=False), UI.STATUS_COLOURS)
        # Prism pins a light palette, so the screen paints the light set.
        self.assertEqual(UI.status_colours(), UI.STATUS_COLOURS)

    def test_a_reply_cell_wears_the_colour_of_the_status_it_leads_to(self):
        self.assertEqual(UI.INTENT_COLOURS["accepted"],
                         UI.STATUS_COLOURS[register.ACCEPTED])
        self.assertEqual(UI.INTENT_COLOURS["rejected"],
                         UI.STATUS_COLOURS[register.NOT_CONVERTED])
        self.assertEqual(UI.INTENT_COLOURS["negotiating"],
                         UI.STATUS_COLOURS[register.NEGOTIATING])


class TheScreenAndExcelReadAStatusAlike(unittest.TestCase):
    """Review F10: a blank status was blue in Excel and untinted on screen,
    and "quoted" or "Won" from an imported sheet was tinted by neither."""

    def test_spellings_from_hand_kept_sheets(self):
        from addons.inquiry import documents
        for typed, status in (("", register.NEW), ("  ", register.NEW),
                              ("quoted", register.QUOTED),
                              ("FOLLOWING UP", register.FOLLOWING_UP),
                              ("Won", register.CONVERTED),
                              ("lost", register.NOT_CONVERTED),
                              ("Converted", register.CONVERTED),
                              ("On hold", "On hold")):
            self.assertEqual(documents.status_key(typed), status, repr(typed))

    def test_each_row_is_the_same_colour_on_screen_and_in_excel(self):
        from openpyxl import load_workbook
        from PySide6.QtCore import Qt
        from PySide6.QtWidgets import QTableWidget
        typed = ["", "quoted", "Won", "lost", "Negotiating", "On hold"]
        rows = [dict(register.blank_row(), **{"Inquiry no": f"INQ/26-27/{n:04d}",
                                              "Status": status})
                for n, status in enumerate(typed, 1)]
        path = documents.register_xlsx(rows, os.path.join(tempfile.mkdtemp(), "r.xlsx"))
        sheet = load_workbook(path).active
        table = QTableWidget(0, 1)
        for n, status in enumerate(typed):
            UI.InquiryDialog._fill_cells(table, [status], status_column=0)
            painted = table.item(n, 0).background()
            fill = sheet.cell(n + 2, 1).fill
            if status == "On hold":
                self.assertEqual(painted.style(), Qt.NoBrush, status)
                self.assertNotEqual(fill.fill_type, "solid", status)
                continue
            screen = painted.color().name().upper().lstrip("#")
            self.assertEqual(fill.fgColor.rgb, "FF" + screen, repr(status))


class APickedRowShowsOnEveryTint(unittest.TestCase):
    """A picked row has to read as picked on each of the twelve tints, and
    no tint may read as picked. Round 1's firmer wash over round 2's strong
    blue did not: a picked row is now a charcoal fill with the tint as its
    ink and a bar in the tint down the left edge (13 Oct, round 3)."""

    def _paint(self, background, selected, tinted=True):
        from PySide6.QtGui import QBrush, QColor, QImage, QPainter
        from PySide6.QtWidgets import (QStyle, QStyleOptionViewItem,
                                       QTableWidget, QTableWidgetItem)
        table = QTableWidget(1, 2)
        item = QTableWidgetItem("")
        if tinted:
            item.setBackground(QBrush(QColor(background)))
        table.setItem(0, 0, item)
        option = QStyleOptionViewItem()
        option.rect.setRect(0, 0, 60, 20)
        if selected:
            option.state |= QStyle.State_Selected
        image = QImage(60, 20, QImage.Format_ARGB32)
        image.fill(QColor("white"))
        painter = QPainter(image)
        UI._TintedCellDelegate(table).paint(painter, option,
                                            table.model().index(0, 0))
        painter.end()
        return image

    @staticmethod
    def _at(image, x, y):
        from PySide6.QtGui import QColor
        colour = QColor(image.pixel(x, y))
        return colour.red(), colour.green(), colour.blue()

    def test_a_new_row_never_looks_picked(self):
        """13 Oct, the real window: every New row on To quote looked
        selected. The New tint is far from the palette's Highlight and from
        how a picked row is painted — over any tint, or over none."""
        from PySide6.QtGui import QPalette
        import theme as T
        for mode in MODES:
            new = theme.REGISTER_TONES[mode]["new"][0]
            if mode == "light":
                highlight = _app.palette().color(QPalette.Highlight)
                rgb = (highlight.red(), highlight.green(), highlight.blue())
                self.assertGreaterEqual(delta_e(new, rgb), 15,
                                        f"New {new} vs Highlight")
                self.assertGreaterEqual(delta_e(new, T.ACCENT), 15)
            unders = list(theme.REGISTER_TONES[mode].values()) + [(None, None)]
            for bg, _ink in unders:
                picked = self._at(self._paint(bg or "#ffffff", True,
                                              tinted=bg is not None), 40, 10)
                d = delta_e(new, picked)
                self.assertGreaterEqual(d, 15, f"{mode} New {new} vs a picked "
                                        f"row over {bg}: ΔE00 {d:.1f}")

    def test_a_picked_row_is_readable(self):
        """The picked fill is charcoal with the tint as its ink, so a picked
        row keeps its state's colour without losing its words."""
        for mode in MODES:
            for state, (bg, _ink) in theme.REGISTER_TONES[mode].items():
                fill, ink, _bar = UI._TintedCellDelegate.picked_colours(
                    theme.qcolor(bg))
                ratio = contrast(fill.name(), ink.name())
                self.assertGreaterEqual(ratio, 4.5, f"{mode} {state}: {ratio:.1f}")

    def test_every_tint_in_both_themes(self):
        for mode in MODES:
            for state, (bg, _ink) in theme.REGISTER_TONES[mode].items():
                plain, picked = self._paint(bg, False), self._paint(bg, True)
                wash = delta_e(self._at(plain, 40, 10), self._at(picked, 40, 10))
                self.assertGreaterEqual(wash, 5, f"{mode} {state}: wash {wash:.1f}")
                bar = contrast(self._at(picked, 1, 10), self._at(picked, 40, 10))
                self.assertGreaterEqual(bar, 3, f"{mode} {state}: bar {bar:.1f}")


# ── the screen ────────────────────────────────────────────────────────────
def _message(subject, sender):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = "sales@acme.co.in"
    msg["Date"] = "Mon, 10 Aug 2026 09:14:00 +0530"
    msg["Message-ID"] = f"<{abs(hash(subject + sender))}@x>"
    msg.set_content("Kindly quote 5000 nos compression spring 2mm.")
    return inbox.parse_message(msg.as_bytes(), uid=1)


def _cfg(folder: str) -> dict:
    return {"api_key": "", "inquiry": {
        "account": {"address": "sales@acme.co.in", "password": "p",
                    "host": "mail.acme.co.in", "port": 993, "folder": "INBOX"},
        "folder": folder, "rate_list": "", "cost_sheet": "",
        "company": "Acme Springs", "signature": "Sales",
        "terms": {"gst_percent": 18, "validity_days": 15,
                  "payment": "advance", "delivery": "3 weeks"},
        "followup_days": 3, "local_only": True,
        "knowledge": {"own_domains": ["acme.co.in"],
                      "customers": ["shaktiauto.in"], "vendors": [],
                      "learned": {}},
        "state": {}}}


def _rows(statuses: list[str]) -> list[dict]:
    rows = []
    for n, status in enumerate(statuses, 1):
        row = register.from_message(
            _message(f"Enquiry {n}", f"Buyer {n} <buyer{n}@shaktiauto.in>"))
        row["Inquiry no"] = f"INQ/26-27/{n:04d}"
        if status != register.NEW:
            register.mark_quoted(row, f"QTN/26-27/{n:04d}", "1000")
            row["Status"] = status
        rows.append(row)
    return rows


class ToQuoteKeepsTheQuotedOnesGrey(unittest.TestCase):

    def setUp(self):
        self.folder = tempfile.mkdtemp()
        self.dialog = UI.InquiryDialog(_cfg(self.folder))

    def _load(self, statuses):
        rows = _rows(statuses)
        register.save(rows, mailflow.Paths(self.folder).register_csv)
        self.dialog._refresh_register()
        return rows

    def _colour(self, table, index):
        return table.item(index, 0).background().color()

    def test_a_quoted_row_stays_grey_under_the_ones_still_to_quote(self):
        self._load([register.QUOTED, register.NEW, register.FOLLOWING_UP,
                    register.NEW])
        table = self.dialog.to_quote_table
        self.assertEqual(table.rowCount(), 4)
        order = [table.item(i, 0).text() for i in range(4)]
        self.assertEqual(order, ["INQ/26-27/0002", "INQ/26-27/0004",
                                 "INQ/26-27/0001", "INQ/26-27/0003"])
        blue = theme.qcolor(UI.STATUS_COLOURS[register.NEW][0])
        grey = theme.qcolor(UI.STATUS_COLOURS[register.QUOTED][0])
        self.assertEqual([self._colour(table, i) for i in range(4)],
                         [blue, blue, grey, grey])
        # Whole row, every column, and the word still says what it is.
        for column in range(table.columnCount()):
            self.assertEqual(table.item(3, column).background().color(), grey)
        self.assertEqual(table.item(3, 5).text(), register.FOLLOWING_UP)
        # The table's row list is the table's order, so picking works.
        self.assertEqual([r["Inquiry no"] for r in self.dialog._to_quote_rows],
                         order)

    def test_the_quoted_rows_are_in_no_answer_yet_as_well(self):
        self._load([register.QUOTED, register.FOLLOWING_UP])
        self.assertEqual(self.dialog.followups.rowCount(), 2)
        self.assertEqual(self.dialog.to_quote_table.rowCount(), 2)
        # There, each in its own state's colour: quoted grey, chased amber.
        self.assertEqual(self._colour(self.dialog.followups, 0),
                         theme.qcolor(UI.STATUS_COLOURS[register.QUOTED][0]))
        self.assertEqual(self._colour(self.dialog.followups, 1),
                         theme.qcolor(UI.STATUS_COLOURS[register.FOLLOWING_UP][0]))

    def test_grey_rows_never_count_as_to_quote(self):
        rows = self._load([register.QUOTED, register.NEW, register.FOLLOWING_UP])
        tab = self.dialog.tabs.tabText(UI.TAB_INDEX["to_quote"])
        self.assertTrue(tab.endswith("(1)"), tab)
        self.assertEqual(DATA.needs_you(_cfg(self.folder), rows)["to_quote"], 1)

    def test_only_grey_rows_count_zero(self):
        self._load([register.QUOTED, register.QUOTED])
        self.assertEqual(self.dialog.to_quote_table.rowCount(), 2)
        self.assertTrue(self.dialog.tabs.tabText(
            UI.TAB_INDEX["to_quote"]).endswith("(0)"))

    def test_once_the_customer_answers_it_moves_on(self):
        self._load([register.NEGOTIATING, register.ACCEPTED,
                    register.CONVERTED, register.NOT_CONVERTED])
        self.assertEqual(self.dialog.to_quote_table.rowCount(), 0)

    def test_a_grey_row_offers_a_reminder_not_a_second_quotation(self):
        self._load([register.QUOTED])
        self.dialog.to_quote_table.setCurrentCell(0, 0)
        self.dialog._sync_panel("to_quote")
        row = self.dialog._current_row()
        self.assertEqual(row["Status"], register.QUOTED)
        keys = UI.actions_for(row["Status"], "to_quote")
        self.assertNotIn("prepare", keys)
        self.assertIn("remind", keys)

    def test_all_inquiries_paints_each_state_its_own_colour(self):
        statuses = list(register.STATUSES)
        self._load(statuses)
        table = self.dialog.register_table
        seen = {}
        for index in range(table.rowCount()):
            status = table.item(index, 4).text()
            seen[status] = table.item(index, 0).background().color().name()
        for status in statuses:
            self.assertEqual(seen[status], UI.STATUS_COLOURS[status][0], status)
        self.assertEqual(len({seen[s] for s in statuses}),
                         len(set(UI.STATUS_STATES.values())))


class TheExcelCopyIsColouredTheSame(unittest.TestCase):

    def _book(self, rows):
        from openpyxl import load_workbook
        path = documents.register_xlsx(rows, Path(tempfile.mkdtemp()) / "r.xlsx")
        return load_workbook(path)

    def test_every_row_is_filled_whole_with_its_state(self):
        rows = _rows(list(register.STATUSES))
        sheet = self._book(rows).worksheets[0]
        header = [c.value for c in sheet[1]]
        status_at = header.index("Status")
        self.assertEqual(sheet.max_row, len(rows) + 1)
        for line in sheet.iter_rows(min_row=2):
            status = line[status_at].value
            bg, ink = UI.STATUS_COLOURS[status]
            for cell in line:
                self.assertEqual(cell.fill.fgColor.rgb, "FF" + bg[1:].upper(),
                                 f"{status} {cell.coordinate}")
                self.assertEqual(cell.font.color.rgb, "FF" + ink[1:].upper())
        # The header stays plain and bold.
        self.assertTrue(sheet["A1"].font.bold)
        self.assertNotEqual(sheet["A1"].fill.fill_type, "solid")

    def test_a_row_with_no_status_is_new(self):
        row = register.blank_row()
        row.update({"Inquiry no": "X", "Status": ""})
        sheet = self._book([row]).worksheets[0]
        self.assertEqual(sheet["A2"].fill.fgColor.rgb,
                         "FF" + UI.STATUS_COLOURS[register.NEW][0][1:].upper())

    def test_the_colour_key_matches_the_rows(self):
        book = self._book(_rows([register.NEW]))
        self.assertEqual(len(book.worksheets), 2)
        key = book.worksheets[1]
        listed = {line[0].value: line[0].fill.fgColor.rgb
                  for line in key.iter_rows(min_row=2)}
        self.assertEqual(set(listed), set(register.STATUSES))
        for status, rgb in listed.items():
            self.assertEqual(rgb, "FF" + UI.STATUS_COLOURS[status][0][1:].upper())
        # Every meaning is written out, not left to the colour alone.
        self.assertTrue(all(line[1].value for line in key.iter_rows(min_row=2)))


if __name__ == "__main__":
    unittest.main()
