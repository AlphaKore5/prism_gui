"""What an AI agent knows: the notes it has been given, the state of each, and the windows that add more.

A note is `ready` (the agent can use it), `processing` (a website or a page being read — this list checks again every few seconds until it is done) or `failed` (with
the reason, and a button to try again). A web page can be read again when the page changed. Notes that can also be found by what they MEAN, not only by their words,
say so. Adding is by pasting text, a question and its answer, one web page, a whole website, or a file.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QFileDialog, QHBoxLayout, QLineEdit, QPlainTextEdit, QSpinBox, QVBoxLayout, QWidget

import i18n
import theme
from addons.whatsapp import ui
from dialogs.base import PrismDialog
from widgets import controls as C

MAX_FILE_BYTES = 6 * 1024 * 1024
TEXT_SUFFIXES = (".txt", ".md", ".markdown", ".csv", ".json", ".log")
POLL_MS = 3000


def read_document(path: str) -> str:
    """The text of a file an owner picked, or ValueError saying why it cannot be read here."""
    p = Path(path)
    if not p.is_file():
        raise ValueError(i18n.t("That file is not there."))
    if p.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(i18n.t("That file is larger than 6 MB. Split it, or paste the part that matters."))
    suffix = p.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        from addons.whatsapp.import_dialog import decode_text
        text = decode_text(p.read_bytes())
        if suffix == ".csv":                                           # a table reads better as lines than as commas
            text = "\n".join(" · ".join(c.strip() for c in row if c.strip()) for row in csv.reader(io.StringIO(text)))
        return text
    if suffix == ".docx":
        try:
            import docx
        except ImportError:
            raise ValueError(i18n.t("Reading Word files is not available in this copy of Prism. Paste the text instead.")) from None
        try:
            d = docx.Document(str(p))
        except Exception as e:                                          # noqa: BLE001
            raise ValueError(i18n.t("That Word file could not be opened ({e}).").format(e=e)) from None
        lines = [para.text for para in d.paragraphs if para.text.strip()]
        for table in d.tables:
            for row in table.rows:
                lines.append(" · ".join(c.text.strip() for c in row.cells if c.text.strip()))
        return "\n\n".join(lines)
    if suffix == ".pdf":
        for module in ("pypdf", "PyPDF2"):
            try:
                reader = __import__(module).PdfReader(str(p))
            except ImportError:
                continue
            except Exception as e:                                      # noqa: BLE001
                raise ValueError(i18n.t("That PDF could not be opened ({e}).").format(e=e)) from None
            text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
            if not text.strip():
                raise ValueError(i18n.t("That PDF has no text to read — it may be a scan. Paste the text instead."))
            return text
    raise ValueError(i18n.t("Prism can read .txt, .md, .csv, .docx and .pdf files here. For anything else, paste the text."))


class KnowledgeDialog(PrismDialog):
    """Teach the agent one thing: text, a question and its answer, a web page, a whole website, or a file."""

    def __init__(self, agent_id: str, mode: str, parent=None):
        titles = {"text": i18n.t("Write what it should know"), "faq": i18n.t("A question and its answer"), "url": i18n.t("Read a web page"),
                  "site": i18n.t("Read a whole website"), "file": i18n.t("Teach it from a file")}
        super().__init__(titles[mode], i18n.t("The agent answers from this. Keep it true and current."), icon="file-text", parent=parent)
        self.agent_id, self.mode = agent_id, mode
        self.title = QLineEdit()
        self.title.setMinimumHeight(38)
        self.title.setPlaceholderText(i18n.t("A short title, e.g. Delivery"))
        self.text = QPlainTextEdit()
        self.text.setMinimumHeight(180)
        self.answer = QPlainTextEdit()
        self.answer.setFixedHeight(120)
        self.url = QLineEdit()
        self.url.setMinimumHeight(38)
        self.url.setPlaceholderText("https://…")
        self.pages = QSpinBox()
        self.pages.setRange(1, 30)
        self.pages.setValue(15)
        self.pages.setMinimumHeight(36)
        self.file_label = C.label(i18n.t("No file chosen"), level="SUPPORT", colour=theme.NEUTRAL[600], wrap=True)
        self.path = ""
        if mode == "text":
            self.body.addWidget(ui.field(i18n.t("Title"), self.title))
            self.text.setPlaceholderText(i18n.t("Paste or write it here — prices, timings, how you work, what you do not do…"))
            self.body.addWidget(ui.field(i18n.t("What it should know"), self.text))
        elif mode == "faq":
            self.title.setPlaceholderText(i18n.t("The question a customer asks"))
            self.body.addWidget(ui.field(i18n.t("Question"), self.title))
            self.answer.setPlaceholderText(i18n.t("The answer you want given"))
            self.body.addWidget(ui.field(i18n.t("Answer"), self.answer))
        elif mode == "url":
            self.body.addWidget(ui.field(i18n.t("Web address"), self.url, i18n.t("Only the one page is read, not the whole site. It must start with https://")))
            self.body.addWidget(ui.field(i18n.t("Title (optional)"), self.title))
        elif mode == "site":
            self.body.addWidget(ui.field(i18n.t("Website"), self.url, i18n.t("Start with its front page. Prism reads that page and the pages it links to on the same site, "
                                                                              "up to two clicks away. It must start with https://")))
            self.body.addWidget(ui.field(i18n.t("Read at most this many pages"), self.pages, i18n.t("It happens in the background; each page shows up in the list when it is read.")))
        else:
            row = QHBoxLayout()
            row.addWidget(C.button(i18n.t("Choose file…"), "secondary", on_click=self._choose))
            row.addWidget(self.file_label, 1)
            self.body.addLayout(row)
            self.body.addWidget(C.label(i18n.t("Text, Markdown, CSV, Word and PDF files up to 6 MB."), level="META", colour=theme.NEUTRAL[500]))
            self.body.addWidget(ui.field(i18n.t("Title"), self.title))
        self.error = ui.note()
        self.body.addWidget(self.error)
        self.body.addStretch(1)
        self.go = self.button(i18n.t("Teach it"), "primary", on_click=self._go)
        self.footer.add_secondary(self.button(i18n.t("Cancel"), on_click=self.reject))
        self.footer.set_primary(self.go)
        self.resize(560, 560 if mode in ("text", "file") else 480)

    def _choose(self):
        path, _ = QFileDialog.getOpenFileName(self, i18n.t("Choose a file"), "", i18n.t("Documents (*.txt *.md *.csv *.docx *.pdf);;All files (*)"))
        if path:
            self.load(path)

    def load(self, path: str):
        try:
            self.text.setPlainText(read_document(path))
        except ValueError as e:
            self.error.error(str(e))
            return
        self.error.clear_note()
        self.path = path
        self.file_label.setText("%s — %s" % (Path(path).name, i18n.t("{n} characters read").format(n=len(self.text.toPlainText()))))
        if not self.title.text().strip():
            self.title.setText(Path(path).stem[:120])

    def _go(self):
        m = self.mode
        if m == "faq":
            route, data = "ai/knowledge/add_faq", {"agent_id": self.agent_id, "question": self.title.text().strip(), "answer": self.answer.toPlainText().strip()}
        elif m == "url":
            route, data = "ai/knowledge/add_url", {"agent_id": self.agent_id, "url": self.url.text().strip(), "title": self.title.text().strip()}
        elif m == "site":
            route, data = "ai/knowledge/add_site", {"agent_id": self.agent_id, "url": self.url.text().strip(), "max_pages": self.pages.value()}
        else:
            body = self.text.toPlainText().strip()
            if m == "file" and not body:
                self.error.error(i18n.t("Choose a file first."))
                return
            route, data = "ai/knowledge/add_text", {"agent_id": self.agent_id, "title": self.title.text().strip(), "body": body,
                                                    "kind": "file" if m == "file" else "text", "source": Path(self.path).name if self.path else ""}
        self.go.setEnabled(False)
        ui.call(self, route, data, done=lambda r: self.accept(), failed=self._failed)

    def _failed(self, message: str, _detail: dict):
        self.go.setEnabled(True)
        self.error.error(message)


def kind_words(kind: str) -> str:
    return {"text": i18n.t("Text"), "faq": i18n.t("Q & A"), "url": i18n.t("Web page"), "file": i18n.t("File")}.get(kind, kind or "")


def state_words(doc: dict) -> tuple:
    """(words, tone) for a source's state."""
    s = doc.get("status") or "ready"
    return {"ready": (i18n.t("Ready"), "ok"), "processing": (i18n.t("Reading…"), "info"), "failed": (i18n.t("Couldn't read"), "err")}.get(s, (s, "quiet"))


class KnowledgeSection(QWidget):
    """The list of what the agent knows, with the buttons that add to it."""

    def __init__(self, editor, parent=None):
        super().__init__(parent)
        self.editor = editor
        self.docs: list = []
        col = QVBoxLayout(self)
        col.setContentsMargins(0, 0, 0, 0)
        card, kc = ui.card(i18n.t("What it knows"), i18n.t("Everything it says comes from here, so add what customers ask about."))
        self.kb_col = QVBoxLayout()
        self.kb_col.setSpacing(theme.SPACE_2)
        kc.addLayout(self.kb_col)
        add = QHBoxLayout()
        add.setSpacing(theme.SPACE_1)
        for mode, text in (("text", i18n.t("Write text")), ("faq", i18n.t("Question and answer")), ("url", i18n.t("Web page")), ("site", i18n.t("Whole website")), ("file", i18n.t("File"))):
            add.addWidget(C.button(text, "secondary", on_click=lambda _=False, m=mode: self.teach(m)))
        add.addStretch(1)
        kc.addLayout(add)
        col.addWidget(card)
        col.addStretch(1)
        self.timer = QTimer(self)
        self.timer.setInterval(POLL_MS)
        self.timer.timeout.connect(self.load)

    def load(self):
        ui.call(self.editor, "ai/knowledge/list", {"agent_id": self.editor.agent["id"]}, done=self._loaded, failed=lambda m, _d: self.editor.say(m, error=True))

    def stop(self):
        self.timer.stop()

    def _loaded(self, reply: dict):
        self.docs = list(reply.get("knowledge") or [])
        ui.clear(self.kb_col)
        if not self.docs:
            self.kb_col.addWidget(C.label(i18n.t("It has not been taught anything yet."), level="SUPPORT", colour=theme.NEUTRAL[600]))
        for d in self.docs:
            self.kb_col.addWidget(self._row(d))
        busy = any(d.get("status") == "processing" for d in self.docs)
        if busy and not self.timer.isActive():
            self.timer.start()
        elif not busy:
            self.timer.stop()

    def _row(self, d: dict) -> QWidget:
        row = QHBoxLayout()
        row.setSpacing(theme.SPACE_2)
        names = QVBoxLayout()
        names.setSpacing(0)
        names.addWidget(C.label(d.get("title") or "", level="SUPPORT", weight=600, wrap=True))
        if d.get("source") and d.get("kind") == "url":
            names.addWidget(C.label(d["source"], level="META", colour=theme.NEUTRAL[500]))
        row.addLayout(names, 1)
        row.addWidget(C.Pill(kind_words(d.get("kind")), "neutral"))
        words, tone = state_words(d)
        pill = C.Pill(words, tone)
        if d.get("error"):
            pill.setToolTip(d["error"])
        row.addWidget(pill)
        if d.get("chunks") and d.get("embedded") == d.get("chunks"):
            by_meaning = C.Pill(i18n.t("Found by meaning"), "quiet")
            by_meaning.setToolTip(i18n.t("A customer who words a question differently from this note can still be answered from it."))
            row.addWidget(by_meaning)
        row.addWidget(C.label(i18n.t("{n} characters").format(n="{:,}".format(int(d.get("characters") or 0))), level="META", colour=theme.NEUTRAL[500]))
        if d.get("kind") == "url" and d.get("status") != "processing":
            row.addWidget(C.button(i18n.t("Read again"), "tertiary", small=True, on_click=lambda _=False, x=d: self.refresh(x)))
        row.addWidget(C.icon_button("trash", i18n.t("Forget this"), on_click=lambda _=False, x=d: self.forget(x), colour=theme.NEUTRAL[600]))
        box = QWidget()
        box.setLayout(row)
        return box

    def teach(self, mode: str):
        if KnowledgeDialog(self.editor.agent["id"], mode, self.editor).exec():
            self.editor.say(i18n.t("Learned.") if mode != "site" else i18n.t("Reading the website. Its pages appear below as they are read."))
            self.load()
            self.editor.owner.refresh_list()

    def forget(self, doc: dict):
        if ui.confirm(self.editor, i18n.t("Forget this"), i18n.t("Make the agent forget “{t}”?").format(t=doc.get("title"))):
            ui.call(self.editor, "ai/knowledge/delete", {"knowledge_id": doc["id"]}, done=lambda r: (self.load(), self.editor.owner.refresh_list()))

    def refresh(self, doc: dict):
        ui.call(self.editor, "ai/knowledge/refresh", {"knowledge_id": doc["id"]}, done=lambda r: (self.editor.say(i18n.t("Reading it again.")), self.load()),
                failed=lambda m, _d: self.editor.say(m, error=True))
