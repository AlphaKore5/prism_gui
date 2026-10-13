"""
The files the inquiry screen hands to other people
──────────────────────────────────────────────────
[3] "Proper formatting and saving of files". What a customer received was a
bare CSV (raw 1100000.00, no letterhead) and a fixed-width text table that
falls apart in Gmail and Outlook, which use proportional fonts. What the
owner opened in Excel was a CSV whose dd-mm-yyyy dates Excel re-reads by the
machine's locale. This module makes the two documents properly:

  · quotation_pdf()   — the quotation as an A4 PDF on the company's name,
                        ₹ in lakh grouping, an aligned table. Qt draws it
                        (QTextDocument + QPdfWriter), so nothing new has to
                        be installed or bundled on Windows or macOS.
  · register_xlsx()   — the register as an Excel workbook with real dates,
                        phone numbers kept as text, money as numbers.
                        openpyxl is already a dependency and in every build
                        (requirements.txt, packaging/prism.spec).

The CSV files stay as they are: the register CSV is the book Prism reads
back, and the quotation CSV is what "Write a reply to win them back" reads
the sent figures from. These are the copies for people.
"""
from __future__ import annotations

import html
import re
from datetime import date
from pathlib import Path

import core_bridge as CB
import i18n


def _money(value) -> str:
    return "₹ " + CB.get_quoting().indian_currency(value)


def quotation_html(quote, settings: dict) -> str:
    """The quotation as one HTML page, every customer-supplied word escaped.

    Kept to the subset QTextDocument draws faithfully: tables with cell
    padding, inline colours, no CSS layout."""
    e = html.escape
    settings = settings or {}
    company = (settings.get("company") or "").strip()
    signature = (settings.get("signature") or "").strip()
    when = quote.date or date.today()
    terms = quote.terms

    head = []
    if company:
        head.append(f'<div style="font-size:18pt; font-weight:bold;">{e(company)}</div>')
    head.append(f'<div style="font-size:14pt; font-weight:bold; margin-top:6px;">'
                f'{e(i18n.t("Quotation"))} {e(quote.number)}</div>')
    meta = [(i18n.t("Date"), when.strftime("%d-%m-%Y")),
            (i18n.t("Valid until"), quote.valid_until.strftime("%d-%m-%Y"))]
    if quote.inquiry_no:
        meta.append((i18n.t("Your inquiry"), quote.inquiry_no))
    to = [quote.customer, quote.contact, quote.email]
    to = [t for t in dict.fromkeys(t.strip() for t in to if t and t.strip())]

    rows = []
    for index, line in enumerate(quote.lines, 1):
        rows.append(
            "<tr>"
            f'<td align="right">{index}</td>'
            f"<td>{e(line.description)}</td>"
            f"<td>{e(line.hsn or '')}</td>"
            f'<td align="right">{e(str(line.quantity))} {e(line.unit or "")}</td>'
            f'<td align="right">{e(_money(line.rate))}</td>'
            f'<td align="right">{e(_money(line.amount))}</td>'
            "</tr>")
    totals = [(i18n.t("Subtotal"), _money(quote.subtotal))]
    if terms.discount_percent:
        totals.append((f"{i18n.t('Discount')} {terms.discount_percent}%",
                       "- " + _money(quote.discount)))
    if terms.freight:
        totals.append((i18n.t("Freight"), _money(terms.freight)))
    totals.append((f"GST {terms.gst_percent}%", _money(quote.gst)))
    for label, value in totals:
        rows.append(f'<tr><td colspan="5" align="right">{e(label)}</td>'
                    f'<td align="right">{e(value)}</td></tr>')
    rows.append(f'<tr><td colspan="5" align="right"><b>{e(i18n.t("Total"))}</b></td>'
                f'<td align="right"><b>{e(_money(quote.total))}</b></td></tr>')

    lines = [f"<p><b>{e(i18n.t('Delivery'))}:</b> {e(terms.delivery)}</p>"
             if terms.delivery else "",
             f"<p><b>{e(i18n.t('Payment'))}:</b> {e(terms.payment)}</p>"
             if terms.payment else ""]
    sign = "<br/>".join(e(part) for part in signature.splitlines()) if signature else ""
    if company:
        sign = (f"{e(i18n.t('For'))} {e(company)}" + (f"<br/>{sign}" if sign else ""))

    header = ("<th>{}</th>" * 6).format(*(e(i18n.t(h)) for h in (
        "Sr", "Description", "HSN", "Quantity", "Rate", "Amount")))
    return (
        '<html><body style="font-family: sans-serif; font-size: 10pt;">'
        + "".join(head)
        + '<table cellspacing="0" cellpadding="2" style="margin-top:8px;">'
        + "".join(f"<tr><td>{e(k)}:</td><td>{e(v)}</td></tr>" for k, v in meta)
        + "</table>"
        + (f'<p style="margin-top:8px;"><b>{e(i18n.t("To"))}:</b><br/>'
           + "<br/>".join(e(t) for t in to) + "</p>" if to else "")
        + '<table width="100%" border="1" cellspacing="0" cellpadding="5" '
          'style="border-collapse:collapse; border-color:#a1a1aa; margin-top:8px;">'
        + f'<tr style="background-color:#e5e5ea;">{header}</tr>'
        + "".join(rows) + "</table>"
        + "".join(lines)
        + (f'<p style="margin-top:18px;">{sign}</p>' if sign else "")
        + "</body></html>")


def plain_summary(quote) -> str:
    """The priced lines for the body of the covering mail, one per line.

    Not a table. Gmail and Outlook show mail in a proportional font, and
    the fixed-width columns render_text() lines up in a terminal came out
    as a ragged mess there; the PDF carries the table."""
    money = _money
    out = []
    for index, line in enumerate(quote.lines, 1):
        out.append(f"{index}. {line.description} — {line.quantity} {line.unit} "
                   f"× {money(line.rate)} = {money(line.amount)}")
    out.append("")
    out.append(f"{i18n.t('Subtotal')}: {money(quote.subtotal)}")
    if quote.terms.discount_percent:
        out.append(f"{i18n.t('Discount')} {quote.terms.discount_percent}%: "
                   f"- {money(quote.discount)}")
    if quote.terms.freight:
        out.append(f"{i18n.t('Freight')}: {money(quote.terms.freight)}")
    out.append(f"GST {quote.terms.gst_percent}%: {money(quote.gst)}")
    out.append(f"{i18n.t('Total')}: {money(quote.total)}")
    return "\n".join(out)


def quotation_pdf(quote, settings: dict, path) -> str:
    """Write the quotation as an A4 PDF at `path`. Returns the path written,
    or "" when it could not be written (the caller then sends the CSV, as
    before, rather than nothing — see pdf_fallback_note()).

    Drawn into a temporary file beside the target and moved over it only
    once it is complete. QPdfWriter does not raise when it cannot open its
    file — it logs "QPainter::begin(): Returned false" and carries on — so
    the old check, "the target exists and is not empty", passed on an older
    PDF of the same name and that stale quotation went to the customer. On
    Windows that is exactly what happens while the PDF is open in Acrobat.
    A file we just made is the only proof we wrote it."""
    from PySide6.QtCore import QMarginsF, QSizeF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

    target = Path(path)
    part = target.with_name(target.name + ".part")
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        _remove(part)
        writer = QPdfWriter(str(part))
        writer.setPageSize(QPageSize(QPageSize.A4))
        writer.setPageMargins(QMarginsF(15, 15, 15, 15), QPageLayout.Millimeter)
        writer.setResolution(96)
        writer.setTitle(f"{i18n.t('Quotation')} {quote.number}")
        writer.setCreator("Prism")
        document = QTextDocument()
        document.setHtml(quotation_html(quote, settings))
        area = writer.pageLayout().paintRectPixels(writer.resolution())
        document.setPageSize(QSizeF(area.width(), area.height()))
        document.print_(writer)
        del writer                       # flushes and closes the file
        if not (part.exists() and part.stat().st_size > 0):
            _remove(part)
            return ""
        _retry_replace(part, target)
    except Exception:                                   # noqa: BLE001
        # PermissionError here is the PDF open in another program (Windows
        # locks it); anything else is a disk or Qt failure. Either way the
        # file at `target`, if any, is not this quotation.
        _remove(part)
        return ""
    return str(target)


def pdf_fallback_note() -> str:
    """What the owner is told when quotation_pdf() returned "": plainly, and
    with the one fix that usually works."""
    return i18n.t("Prism couldn't write the quotation PDF. It may be open in "
                  "another program — close it in Acrobat or your PDF viewer "
                  "and try again. The quotation goes as a CSV file instead.")


def _retry_replace(source: Path, target: Path, attempts: int = 10) -> None:
    """os.replace, retried for about half a second: on Windows a virus
    scanner or the preview pane holds a new file for a moment. A file open
    in Acrobat outlasts the retries and the PermissionError is raised."""
    import os
    import time
    for attempt in range(attempts):
        try:
            os.replace(source, target)
            return
        except PermissionError:
            if attempt == attempts - 1:
                raise
            time.sleep(0.05)


def _remove(path: Path) -> None:
    try:
        path.unlink()
    except OSError:
        pass


def quotation_files(quote, settings: dict, folder) -> tuple[str, str]:
    """(csv, pdf) for one quotation in the inquiry's folder, named by its
    number. The CSV is the record Prism reads back; the PDF is what the
    customer is sent. pdf is "" if it could not be drawn."""
    quoting = CB.get_quoting()
    name = safe_name(quote.number.replace("/", "-"))
    written = quoting.write_csv(quote, str(Path(folder) / f"{name}.csv"))
    return written, quotation_pdf(quote, settings, Path(folder) / f"{name}.pdf")


# Only Prism's own names: its quotations (quoting.next_quote_number() always
# numbers QTN/yy-yy/n, saved as QTN-yy-yy-n) and the .eml copies the mail
# check saves (INQ-26-27-0005_mail_2026-10-12_0914.eml). Any LETTERS-yy-yy-n
# used to count, which hid a customer's own PO-26-27-0012.pdf saved in a
# folder before attachments carried the inquiry-number prefix.
_OURS = re.compile(r"^(QTN-\d{2}-\d{2}-\d+\.(pdf|csv|xlsx)(\.part)?"
                   r"|[\w-]*_mail_(\d{4}-\d{2}-\d{2}_\d{4}|undated)(-\d+)?\.eml)$",
                   re.I)


def is_ours(name: str) -> bool:
    """A file Prism wrote into the inquiry's folder (QTN-26-27-0001.pdf, the
    saved mail), as opposed to something the customer sent. The quotation PDF
    must never be offered as "the drawing" or picked up as "the PO"."""
    return bool(_OURS.match(Path(str(name or "")).name))


def attachment_for(csv_path: str, pdf_path: str) -> dict:
    """The one file the customer gets: the PDF, or the CSV if there is none."""
    path = pdf_path or csv_path
    return {"path": path, "name": Path(path).name,
            "mime": "application/pdf" if pdf_path else "text/csv"}


# ── the register, for Excel ───────────────────────────────────────────────────
_DATE_COLUMNS = ("Date received", "Quotation date", "Last contact", "PO date")
_MONEY_COLUMNS = ("Quotation value", "Order value")
_TEXT_COLUMNS = ("Phone", "Inquiry no", "Quotation no", "PO number")
_HIDDEN = ("Thread",)


def register_xlsx(rows: list[dict], path) -> str:
    """The register as an .xlsx a person can open on any locale.

    Dates are real dates shown dd-mm-yyyy (a CSV's "02-08-2026" is the 8th
    of February to an Excel set to US), phone numbers are text so a leading
    0 survives, money is a number Excel can add up, the header is bold and
    stays put, and the columns are wide enough to read. Each row is filled
    with its status's colour, the same one the screen paints ([12][15][17]),
    and a second sheet is the colour key. Raises what openpyxl or the disk
    raise."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as illegal
    from openpyxl.utils import get_column_letter

    register = CB.get_register()
    columns = [c for c in register.COLUMNS if c not in _HIDDEN]
    columns += [k for row in rows for k in row
                if k and k not in columns and k not in _HIDDEN]
    columns = list(dict.fromkeys(columns))

    book = Workbook()
    sheet = book.active
    sheet.title = i18n.t("Inquiries")[:31]
    sheet.append([illegal.sub("", i18n.t(c)) for c in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        _as_text(cell)
    fills = status_fills()
    for row in rows:
        values = []
        for column in columns:
            # Control characters (OCR text, some subjects) make openpyxl
            # refuse the whole workbook; they mean nothing in a cell.
            raw = illegal.sub("", str(row.get(column) or "")).strip()
            if column in _DATE_COLUMNS:
                values.append(register.parse_date(raw) or raw or None)
            elif column in _MONEY_COLUMNS and raw:
                values.append(float(register.money(raw)))
            else:
                values.append(raw or None)
        sheet.append(values)
        line = sheet.max_row
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(line, index)
            if column in _DATE_COLUMNS and isinstance(cell.value, date):
                cell.number_format = "DD-MM-YYYY"
            elif column in _MONEY_COLUMNS and isinstance(cell.value, float):
                cell.number_format = "[$₹-4009] #,##,##0.00"
            elif column in _TEXT_COLUMNS:
                cell.number_format = "@"
            _as_text(cell)
        status = status_key(row.get("Status"))
        if status in fills:
            fill = PatternFill("solid", fgColor=fills[status][0])
            ink = Font(color=fills[status][1])
            for index in range(1, len(columns) + 1):
                sheet.cell(line, index).fill = fill
                sheet.cell(line, index).font = ink
    for index, column in enumerate(columns, 1):
        longest = max([len(str(column))] + [len(str(r.get(column) or ""))
                                            for r in rows[:500]])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    _legend_sheet(book, fills)
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(target))
    return str(target)


# What a spreadsheet reads as the start of a formula. Customer, Product asked
# and Notes come from other people's mail.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _as_text(cell) -> None:
    """Keep a cell that looks like a formula a plain string. openpyxl stores
    any string starting with "=" as a live formula, so a mail subject
    '=HYPERLINK("http://…","Click")' became a link in the owner's Excel.
    quotePrefix also keeps it text when someone edits the cell."""
    value = cell.value
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        cell.data_type = "s"
        cell.quotePrefix = True


# Status spellings from hand-kept or imported sheets, read as Prism's own.
_STATUS_ALIASES = {
    "won": "Converted", "order": "Converted", "order received": "Converted",
    "lost": "Not converted", "not converted": "Not converted",
    "follow up": "Following up", "followup": "Following up",
    "following-up": "Following up", "negotiation": "Negotiating",
    "quote sent": "Quoted", "quotation sent": "Quoted",
}


def status_key(text) -> str:
    """The register status a cell means, for colouring: Prism's own spelling
    in any case, a few common hand-typed words, and blank as New (an inquiry
    nobody has touched — the engine reads it the same way). Anything else is
    returned as typed and gets no colour. One function for the screen and
    the Excel copy, so the two always tint a row alike."""
    raw = " ".join(str(text or "").split())
    if not raw:
        return "New"
    folded = raw.casefold()
    for status in ("New", "Quoted", "Following up", "Negotiating", "Accepted",
                   "Converted", "Not converted"):
        if folded == status.casefold():
            return status
    return _STATUS_ALIASES.get(folded, raw)


def legend() -> list[tuple[str, str]]:
    """What each colour in the register means, in the owner's words, as
    (status, meaning) with the meaning translated. The order is the
    pipeline's, so the key reads top to bottom as an inquiry's life."""
    return [
        ("New", i18n.t("Yet to quote")),
        ("Quoted", i18n.t("Quotation sent, waiting for their answer")),
        ("Following up", i18n.t("No reply yet, reminder sent")),
        ("Negotiating", i18n.t("They answered and want a better rate or "
                               "more detail")),
        ("Accepted", i18n.t("They said yes, PO still to be recorded")),
        ("Converted", i18n.t("Order won, PO received")),
        ("Not converted", i18n.t("Lost")),
    ]


def status_fills() -> dict:
    """{register status: (background, ink)} as openpyxl's ARGB, from the
    screen's own colours (dialog.STATUS_COLOURS, theme.REGISTER_TONES).
    Always the light set: a spreadsheet is printed and read on white."""
    try:
        from addons.inquiry.dialog import STATUS_COLOURS
    except Exception:                                   # noqa: BLE001
        return {}
    argb = lambda value: "FF" + value.lstrip("#").upper()
    return {status: (argb(bg), argb(ink))
            for status, (bg, ink) in STATUS_COLOURS.items()}


def _legend_sheet(book, fills: dict) -> None:
    """A "Colour key" sheet: each status in its own colour, and what it
    means. Built from the same fills the rows use, so the two cannot
    disagree."""
    from openpyxl.styles import Font, PatternFill

    sheet = book.create_sheet(i18n.t("Colour key")[:31])
    sheet.append([i18n.t("Status"), i18n.t("Means")])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for status, meaning in legend():
        sheet.append([i18n.t(status), meaning])
        if status in fills:
            bg, ink = fills[status]
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor=bg)
                cell.font = Font(color=ink)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 58


# ── names that are safe on Windows and macOS ─────────────────────────────────
_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)),
             *(f"LPT{i}" for i in range(1, 10))}


def safe_name(text: str, limit: int = 80) -> str:
    """A file-name part Windows, macOS and Linux all accept: no \\ / : * ? " <
    > | or control characters, no trailing dot or space, not a reserved
    device name, and short enough to leave room for the folder path under
    Windows' 260-character limit."""
    cleaned = "".join("-" if (ch in '\\/:*?"<>|' or ord(ch) < 32) else ch
                      for ch in str(text or ""))
    cleaned = " ".join(cleaned.split()).strip(" .")[:limit].strip(" .")
    if cleaned.split(".")[0].upper() in _RESERVED:
        cleaned = f"_{cleaned}"
    return cleaned
